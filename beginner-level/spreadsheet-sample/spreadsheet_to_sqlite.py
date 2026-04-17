# Read an Excel spreadsheet, convert it to a SQLite table, and insert all rows
#
# This script:
#   1. Opens an Excel file (.xlsx)
#   2. Reads the header row to determine column names
#   3. Detects column types from the data (text, integer, real)
#   4. Creates a SQLite database and table automatically
#   5. Inserts every row from the spreadsheet into the table
#
# Usage:
#   python spreadsheet_to_sqlite.py
#   python spreadsheet_to_sqlite.py --file my_data.xlsx --db my_database.db

import argparse
import os
import sqlite3

import openpyxl


def detect_column_type(values):
    """Detect the SQLite column type based on sample values."""
    has_int = False
    has_float = False

    for val in values:
        if val is None:
            continue
        if isinstance(val, bool):
            has_int = True
        elif isinstance(val, int):
            has_int = True
        elif isinstance(val, float):
            has_float = True
        else:
            return "TEXT"

    if has_float:
        return "REAL"
    if has_int:
        return "INTEGER"
    return "TEXT"


def read_spreadsheet(filepath):
    """Read an Excel file and return (headers, rows, sheet_name)."""
    workbook = openpyxl.load_workbook(filepath, read_only=True)
    sheet = workbook.active
    sheet_name = sheet.title

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("Spreadsheet is empty.")
        return None, None, None

    headers = [str(cell).strip().lower().replace(" ", "_") for cell in rows[0]]
    data_rows = rows[1:]

    print(f"Read '{filepath}' — sheet: '{sheet_name}'")
    print(f"  Columns: {headers}")
    print(f"  Rows: {len(data_rows)}")

    return headers, data_rows, sheet_name


def create_table(cursor, table_name, headers, data_rows):
    """Create a SQLite table based on the spreadsheet structure."""
    # Collect sample values per column to detect types
    column_samples = {h: [] for h in headers}
    for row in data_rows[:100]:
        for i, header in enumerate(headers):
            if i < len(row):
                column_samples[header].append(row[i])

    # Build the CREATE TABLE statement
    column_defs = []
    for header in headers:
        col_type = detect_column_type(column_samples[header])
        column_defs.append(f"    {header} {col_type}")

    create_sql = f"CREATE TABLE IF NOT EXISTS {table_name} (\n"
    create_sql += ",\n".join(column_defs)
    create_sql += "\n);"

    print(f"\nGenerated SQL:\n{create_sql}\n")
    cursor.execute(f"DROP TABLE IF EXISTS {table_name}")
    cursor.execute(create_sql)


def insert_rows(cursor, table_name, headers, data_rows):
    """Insert all rows from the spreadsheet into the SQLite table."""
    placeholders = ", ".join(["?"] * len(headers))
    insert_sql = f"INSERT INTO {table_name} ({', '.join(headers)}) VALUES ({placeholders})"

    inserted = 0
    for row in data_rows:
        # Pad row if it has fewer values than headers
        padded = list(row) + [None] * (len(headers) - len(row))
        # Convert values to strings where needed for TEXT columns
        values = [str(v) if v is not None and not isinstance(v, (int, float)) else v for v in padded[:len(headers)]]
        cursor.execute(insert_sql, values)
        inserted += 1

    print(f"Inserted {inserted} row(s) into '{table_name}'.")
    return inserted


def main():
    parser = argparse.ArgumentParser(
        description="Convert an Excel spreadsheet to a SQLite database table.",
    )
    parser.add_argument(
        "--file", default="employees.xlsx",
        help="Path to the Excel file (default: employees.xlsx)",
    )
    parser.add_argument(
        "--db", default="data.db",
        help="Path to the SQLite database file (default: data.db)",
    )
    args = parser.parse_args()

    if not os.path.exists(args.file):
        print(f"Error: File '{args.file}' not found.")
        print("Run 'python create_sample_spreadsheet.py' first to generate the sample file.")
        return

    # Step 1: Read the spreadsheet
    headers, data_rows, sheet_name = read_spreadsheet(args.file)
    if headers is None:
        return

    # Use the sheet name as the table name
    table_name = sheet_name.lower().replace(" ", "_")

    # Step 2: Create the database and table
    conn = sqlite3.connect(args.db)
    cursor = conn.cursor()

    create_table(cursor, table_name, headers, data_rows)

    # Step 3: Insert all rows
    insert_rows(cursor, table_name, headers, data_rows)

    conn.commit()
    conn.close()

    print(f"\nDone! Database saved to '{args.db}' — table: '{table_name}'")


if __name__ == "__main__":
    main()
